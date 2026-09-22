"""Safe server-side text extraction (Phase 3).

Dispatches on verified kind (never on client MIME). All functions take raw
bytes and enforce output caps. Failures raise ExtractError -> proc failed
status, never a 500 with traceback.
"""
from __future__ import annotations

import csv
import io

MAX_CHARS = 200_000
PDF_MAX_PAGES = 50
SHEET_MAX, ROW_MAX, COL_MAX = 20, 2000, 50


class ExtractError(Exception):
    pass


def _cap(s: str) -> str:
    return s[:MAX_CHARS]


def extract_txt(data: bytes) -> tuple[str, dict]:
    for enc in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return _cap(data.decode(enc)), {"encoding": enc}
        except Exception:
            continue
    raise ExtractError("undecodable text")


def extract_csv(data: bytes) -> tuple[str, dict]:
    text, meta = extract_txt(data)
    try:
        rows = list(csv.reader(io.StringIO(text)))[:ROW_MAX]
    except Exception as e:
        raise ExtractError(f"csv parse: {e}")
    if not rows:
        return "", {**meta, "rows": 0}
    flat = "\n".join(",".join((c or "")[:200] for c in r[:COL_MAX]) for r in rows)
    return _cap(flat), {**meta, "rows": len(rows), "cols": max(len(r) for r in rows)}


def extract_xlsx(data: bytes) -> tuple[str, dict]:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise ExtractError(f"xlsx parse: {e}")
    parts, sheets = [], 0
    try:
        for ws in wb.worksheets[:SHEET_MAX]:
            sheets += 1
            parts.append(f"# {ws.title}")
            for row in ws.iter_rows(values_only=True):
                vals = [(str(v)[:200] if v is not None else "") for v in list(row)[:COL_MAX]]
                if any(vals):
                    parts.append(",".join(vals))
                if len(parts) > ROW_MAX * SHEET_MAX:
                    break
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return _cap("\n".join(parts)), {"sheets": sheets}


def extract_docx(data: bytes) -> tuple[str, dict]:
    try:
        import docx
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise ExtractError(f"docx parse: {e}")
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
    return _cap("\n".join(parts)), {"paragraphs": len(d.paragraphs), "tables": len(d.tables)}


def extract_pdf(data: bytes) -> tuple[str, dict]:
    try:
        from pypdf import PdfReader
        r = PdfReader(io.BytesIO(data), strict=False)  # lenient: real-world PDFs vary
    except Exception as e:
        raise ExtractError(f"pdf parse: {e}")
    if getattr(r, "is_encrypted", False):
        raise ExtractError("encrypted pdf")
    pages = min(len(r.pages), PDF_MAX_PAGES)
    out = []
    for i in range(pages):
        try:
            out.append(r.pages[i].extract_text() or "")
        except Exception:
            out.append("")
    return _cap("\n".join(out)), {"pages": len(r.pages), "pages_read": pages}


def extract(kind: str, data: bytes) -> tuple[str, dict]:
    """kind: txt|csv|xlsx|docx|pdf. Returns (text, meta)."""
    fn = {"txt": extract_txt, "csv": extract_csv, "xlsx": extract_xlsx,
          "docx": extract_docx, "pdf": extract_pdf}.get(kind)
    if fn is None:
        raise ExtractError(f"no text extractor for {kind}")
    text, meta = fn(data)
    return text, {**meta, "chars": len(text)}


def extract_segments(kind: str, data: bytes) -> list[tuple[int | None, str]]:
    """Per-page segments for chunk metadata: [(page_no|None, text)]."""
    if kind == "pdf":
        try:
            from pypdf import PdfReader
            r = PdfReader(io.BytesIO(data), strict=False)
        except Exception as e:
            raise ExtractError(f"pdf parse: {e}")
        if getattr(r, "is_encrypted", False):
            raise ExtractError("encrypted pdf")
        segs = []
        for i in range(min(len(r.pages), PDF_MAX_PAGES)):
            try:
                t = r.pages[i].extract_text() or ""
            except Exception:
                t = ""
            if t.strip():
                segs.append((i + 1, _cap(t)))
        return segs
    text, _ = extract(kind, data)
    return [(None, text)] if text.strip() else []
