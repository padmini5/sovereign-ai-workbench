"""DOCX/XLSX generation with footnote-style citations."""
import os, re
from docx import Document
from . import config as cfg
OUT = cfg.OUTPUTS_DIR
os.makedirs(OUT, exist_ok=True)

def make_docx(title: str, body: str, chunks: list, signed_by: dict, signature: dict) -> str:
    doc = Document()
    doc.add_heading(title, 0)
    for para in body.split("\n"):
        para = para.strip()
        if para:
            doc.add_paragraph(para)
    doc.add_heading("Sources & citations", 1)
    for c in chunks:
        doc.add_paragraph(f"[{c['chunk_id']} p.{c.get('page','?')}] {c.get('title','')} — {c['text'][:200]}",
                          style="List Bullet")
    doc.add_heading("PQC signature", 1)
    doc.add_paragraph(f"Algorithm: {signature.get('algorithm')}")
    doc.add_paragraph(f"Signed by: {signed_by.get('user_id')} ({signed_by.get('role')})")
    doc.add_paragraph(f"SHA-256: {signature.get('payload_sha256')}")
    doc.add_paragraph(f"Signature (b64, truncated): {signature.get('signature_b64','')[:64]}…")
    fn = os.path.join(OUT, re.sub(r"[^A-Za-z0-9_-]+", "_", title)[:40] + "_signed.docx")
    doc.save(fn)
    return fn

def make_xlsx(title: str, rows: list) -> str:
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Findings"
    ws.append(["chunk_id", "page", "text", "score"])
    for c in rows:
        ws.append([c.get("chunk_id"), c.get("page"), c.get("text", "")[:500], c.get("score")])
    fn = os.path.join(OUT, re.sub(r"[^A-Za-z0-9_-]+", "_", title)[:40] + "_evidence.xlsx")
    wb.save(fn)
    return fn
