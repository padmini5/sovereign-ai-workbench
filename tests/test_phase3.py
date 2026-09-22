"""Phase 3 tests: secure document/image upload (run: python tests/test_phase3.py).

Covers: valid uploads (7 types), invalid/mismatched files, oversize,
unauthorized access, ID manipulation, deletion permissions, image
validation, extraction, audit. Isolated temp DBs + temp upload dir.
"""
import io
import os
import sys
import tempfile

_tmp = tempfile.mkdtemp(prefix="sov_phase3_")
os.environ["SOV_AUTH_DB"] = os.path.join(_tmp, "auth.db")
os.environ["SOV_AUTH_AUDIT_DB"] = os.path.join(_tmp, "auth_audit.db")
os.environ["SOV_CHAT_DB"] = os.path.join(_tmp, "chat.db")
os.environ["SOV_DOCS_DB"] = os.path.join(_tmp, "docs.db")
os.environ["SOV_VECTORS_DB"] = os.path.join(_tmp, "vectors.db")
os.environ["SOV_UPLOADS_DIR"] = os.path.join(_tmp, "uploads")
os.environ["SOV_AI_PROVIDER"] = "mock"

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from fastapi.testclient import TestClient  # noqa: E402
from backend.app.main import app  # noqa: E402
from backend.app import audit_log  # noqa: E402

c = TestClient(app)
H = lambda t: {"Authorization": "Bearer " + t}  # noqa: E731

passed = failed = 0
def check(name, cond, extra=""):
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS {name}")
    else:
        failed += 1
        print(f"  FAIL {name} {extra}")


def login(u, p):
    r = c.post("/api/v1/auth/login", json={"username": u, "password": p})
    assert r.status_code == 200, r.text
    return r.json()["access_token"]


def up(tok, name, data, ctype="application/octet-stream"):
    return c.post("/api/v1/docs/upload", files={"f": (name, data, ctype)},
                  headers={"Authorization": f"Bearer {tok}"})


def mk_png():
    from PIL import Image
    im = Image.new("RGB", (64, 32), "red")
    b = io.BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


def mk_jpg():
    from PIL import Image
    im = Image.new("RGB", (64, 32), "blue")
    b = io.BytesIO()
    im.save(b, "JPEG")
    return b.getvalue()


def mk_docx():
    import docx
    d = docx.Document()
    d.add_paragraph("Phase3 docx content alpha beta")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def mk_xlsx():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["col_a", "col_b"])
    ws.append(["gamma", "delta"])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def mk_pdf(text="Hello PDF"):
    """Minimal valid PDF with correct xref/startxref (built programmatically)."""
    stream = f"BT /F1 12 Tf 10 100 Td ({text}) Tj ET\n".encode()
    objs = [b"<</Type/Catalog/Pages 2 0 R>>",
            b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
            b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]/Contents 4 0 R"
            b"/Resources<</Font<</F1 5 0 R>>>>>>",
            b"<</Length %d>>stream\n" % len(stream) + stream + b"endstream",
            b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>"]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj" % i + body + b"endobj\n"
    xref_at = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer<</Size %d/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref_at)
    return bytes(out)


PDF = mk_pdf()

T_OP = login("operator", "Op123!")
T_MGR = login("manager", "Mgr123!")
T_ADM = login("admin", "Admin123!")
T_USR = login("user", "User123!")
T_REV = login("reviewer", "Rev123!")

print("== 1. valid uploads (all 7 types) ==")
samples = [("note.txt", b"hello phase3 world", "text/plain"),
           ("data.csv", b"a,b\n1,2\ngamma,delta\n", "text/csv"),
           ("sheet.xlsx", mk_xlsx(), "application/vnd.ms-excel"),
           ("memo.docx", mk_docx(), "application/msword"),
           ("scan.pdf", PDF, "application/pdf"),
           ("pic.png", mk_png(), "image/png"),
           ("photo.jpg", mk_jpg(), "image/jpeg")]
ids = {}
for name, data, ctype in samples:
    r = up(T_OP, name, data, ctype)
    ok = r.status_code == 200
    check(f"upload {name}", ok, r.text[:200] if not ok else "")
    if ok:
        j = r.json()
        ids[name] = j["id"]
        if name.endswith((".png", ".jpg", ".jpeg")):
            # Sanitized at ingest (EXIF stripped) so byte length may differ;
            # assert identity fields + EXIF-free bytes instead.
            from PIL import Image as _Im
            dl = c.get(f"/api/v1/docs/{j['id']}/download", headers=H(T_OP))
            check(f"{name} fields", j["filename"] == name and j["status"] == "UPLOADED"
                  and j["proc_status"] == "pending" and dl.status_code == 200
                  and _Im.open(io.BytesIO(dl.content)).getexif() == {})
        else:
            check(f"{name} fields", j["filename"] == name and j["size"] == len(data)
                  and j["status"] == "UPLOADED" and j["proc_status"] == "pending")

print("== 2. browser-supplied type is ignored (sniffed) ==")
r = up(T_OP, "pic.png", mk_png(), "text/plain")  # lying content-type
check("wrong client MIME still accepted by bytes", r.status_code == 200, r.text[:150])

print("== 3. invalid files ==")
check(".exe rejected -> 415", up(T_OP, "evil.exe", b"MZ" + b"x" * 100).status_code == 415)
check("no extension -> 415", up(T_OP, "noext", b"hello").status_code == 415)
check("text-as-png -> 415", up(T_OP, "fake.png", b"just text").status_code == 415)
check("pdf-bytes-as-docx -> 415", up(T_OP, "fake.docx", PDF).status_code == 415)
check("corrupt jpg -> 415", up(T_OP, "bad.jpg", b"\xff\xd8\xff" + b"garbage" * 50).status_code == 415)
check("zip-without-office-parts -> 415",
      up(T_OP, "fake.xlsx", b"PK\x03\x04" + b"z" * 200).status_code == 415)
check("empty file -> 400", up(T_OP, "e.txt", b"").status_code == 400)

print("== 4. oversized ==")
os.environ["SOV_MAX_UPLOAD_MB"] = "1"
big = b"a" * (1024 * 1024 + 100)
check("1MB+ file vs 1MB limit -> 413", up(T_OP, "big.txt", big).status_code == 413)
del os.environ["SOV_MAX_UPLOAD_MB"]
check("same file ok at default 10MB", up(T_OP, "big.txt", big).status_code == 200)

print("== 5. unauthorized ==")
check("no token -> 401/403",
      c.post("/api/v1/docs/upload", files={"f": ("a.txt", b"x")}).status_code in (401, 403))
check("USER upload -> 403 (no DOCUMENT_UPLOAD)",
      up(T_USR, "a.txt", b"x").status_code == 403)
check("no-token list -> 401/403", c.get("/api/v1/docs").status_code in (401, 403))
check("OPERATOR analyze -> 403 (no DOCUMENT_ANALYZE)",
      c.post(f"/api/v1/docs/{ids['note.txt']}/analyze", headers=H(T_OP)).status_code == 403)

print("== 6. analyze + extraction (manager: upload+analyze) ==")
r = up(T_MGR, "m.txt", b"manager text omega")
mid = r.json()["id"]
r = c.post(f"/api/v1/docs/{mid}/analyze", headers=H(T_MGR))
check("txt analyze done", r.status_code == 200 and r.json()["proc_status"] == "done", r.text[:200])
check("excerpt contains content", "omega" in r.json().get("excerpt", ""))
r = c.post(f"/api/v1/docs/{ids['memo.docx']}/analyze", headers=H(T_MGR))
check("docx forbidden for non-owner -> 404",
      r.status_code == 404, r.text[:150])
did = up(T_MGR, "m.docx", mk_docx()).json()["id"]
r = c.post(f"/api/v1/docs/{did}/analyze", headers=H(T_MGR)).json()
check("docx text extracted", "alpha beta" in r.get("excerpt", ""), r.get("excerpt", "")[:150])
pid = up(T_MGR, "p.pdf", PDF).json()["id"]
r = c.post(f"/api/v1/docs/{pid}/analyze", headers=H(T_MGR)).json()
check("pdf text extracted", "Hello PDF" in r.get("excerpt", ""), r.get("excerpt", "")[:150])
xid = up(T_MGR, "s.xlsx", mk_xlsx()).json()["id"]
r = c.post(f"/api/v1/docs/{xid}/analyze", headers=H(T_MGR)).json()
check("xlsx cells extracted", "gamma" in r.get("excerpt", ""), r.get("excerpt", "")[:150])

print("== 7. images: preview + metadata-only analyze (no silent OCR) ==")
img_id = up(T_MGR, "im.png", mk_png()).json()["id"]
r = c.post(f"/api/v1/docs/{img_id}/analyze", headers=H(T_MGR)).json()
check("image analyze is metadata-only",
      r["proc_status"] == "metadata_only" and r.get("chunks", -1) == 0,
      str((r.get("proc_status"), r.get("chunks"))))
from backend.app.vectorstore import get_vector_store as _vs
check("no image vectors indexed", _vs().count_for(img_id) == 0)
r = c.get(f"/api/v1/docs/{img_id}/preview", headers=H(T_MGR))
check("thumbnail 200 png", r.status_code == 200 and r.headers["content-type"] == "image/png"
      and r.content.startswith(b"\x89PNG"), r.text[:100] if r.status_code != 200 else "")
d = c.get(f"/api/v1/docs/{img_id}", headers=H(T_MGR)).json()
check("vision metadata stored", d.get("vision", {}).get("width") == 64, str(d.get("vision")))

print("== 8. download integrity ==")
payload = b"download-bytes-check-123"
dl_id = up(T_OP, "dl.txt", payload).json()["id"]
r = c.get(f"/api/v1/docs/{dl_id}/download", headers=H(T_OP))
check("download bytes match", r.status_code == 200 and r.content == payload)

print("== 9. ID manipulation ==")
r = c.post("/api/v1/admin/users",
           json={"username": "tmp_op2", "password": "Tmp12345", "role": "OPERATOR"},
           headers=H(T_ADM))
T_OP2 = login("tmp_op2", "Tmp12345")
victim = ids["note.txt"]
check("stranger detail -> 404", c.get(f"/api/v1/docs/{victim}", headers=H(T_OP2)).status_code == 404)
check("stranger download -> 404",
      c.get(f"/api/v1/docs/{victim}/download", headers=H(T_OP2)).status_code == 404)
check("stranger preview -> 404",
      c.get(f"/api/v1/docs/{victim}/preview", headers=H(T_OP2)).status_code == 404)
check("stranger analyze -> 403/404",
      c.post(f"/api/v1/docs/{victim}/analyze", headers=H(T_OP2)).status_code in (403, 404))
mine = [d["id"] for d in c.get("/api/v1/docs", headers=H(T_OP2)).json()["documents"]]
check("stranger list excludes victim", victim not in mine)
admin_list = [d["id"] for d in c.get("/api/v1/docs", headers=H(T_ADM)).json()["documents"]]
check("ADMIN sees all", victim in admin_list)
check("USER sees only own (empty)", c.get("/api/v1/docs", headers=H(T_USR)).json()["documents"] == [])

print("== 10. deletion permissions ==")
own_del = up(T_OP, "todel.txt", b"bye").json()["id"]
check("owner w/o DOCUMENT_DELETE -> 403",
      c.delete(f"/api/v1/docs/{own_del}", headers=H(T_OP)).status_code == 403)
check("REVIEWER delete -> 403",
      c.delete(f"/api/v1/docs/{own_del}", headers=H(T_REV)).status_code == 403)
check("stranger ADMIN delete other's -> 200",
      c.delete(f"/api/v1/docs/{own_del}", headers=H(T_ADM)).status_code == 200)
check("deleted gone -> 404", c.get(f"/api/v1/docs/{own_del}", headers=H(T_OP)).status_code == 404)
adm_doc = up(T_ADM, "a.txt", b"adm").json()["id"]
check("ADMIN deletes own -> 200",
      c.delete(f"/api/v1/docs/{adm_doc}", headers=H(T_ADM)).status_code == 200)
check("delete missing -> 404",
      c.delete("/api/v1/docs/d-nope", headers=H(T_ADM)).status_code == 404)

print("== 11. search/filter ==")
up(T_OP, "quarterly-report.txt", b"Q1 numbers")
up(T_OP, "quarterly-photo.png", mk_png())
sq = c.get("/api/v1/docs", params={"q": "quarterly"}, headers=H(T_OP)).json()["documents"]
check("search q= filters", len(sq) == 2 and all("quarterly" in d["filename"] for d in sq), str(len(sq)))
si = c.get("/api/v1/docs", params={"kind": "image"}, headers=H(T_OP)).json()["documents"]
check("kind=image filters", len(si) >= 1 and all(d["kind"] == "image" for d in si))
check("bad kind -> 400", c.get("/api/v1/docs", params={"kind": "nope"}, headers=H(T_OP)).status_code == 400)

print("== 12. audit (metadata only) ==")
rows = audit_log.rows(1000)
actions = {r["action"] for r in rows}
for a in ["doc_upload", "doc_view", "doc_download", "doc_analyze", "doc_delete"]:
    check(f"{a} audited", a in actions, str(sorted(actions)))
check("denials audited", "permission_denied" in actions)
blob = " ".join((r.get("detail") or "") for r in rows)
check("no file content in audit", "download-bytes-check-123" not in blob and "omega" not in blob)

# cleanup temp user
uid = [u["id"] for u in c.get("/api/v1/admin/users", headers=H(T_ADM)).json()["users"]
       if u["username"] == "tmp_op2"][0]
c.delete(f"/api/v1/admin/users/{uid}", headers=H(T_ADM))

print(f"\nRESULT: {passed} passed, {failed} failed")
sys.exit(1 if failed else 0)
